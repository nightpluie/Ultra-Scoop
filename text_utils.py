"""ULTRA SCOOP — 文字解析工具：稿件輸出解析、Markdown 清理、附件檔案解析。"""

import os
import re


def _has(m):
    try:
        __import__(m)
        return True
    except ImportError:
        return False


HAS_PDFPLUMBER = _has("pdfplumber")
HAS_DOCX       = _has("docx")
HAS_OPENPYXL   = _has("openpyxl")

# ── output parsers ────────────────────────────────────────────────────────────
_TITLE_RE = re.compile(r'^\[建議標題\]：(.+)', re.MULTILINE)


def _extract_title_body(chunk: str):
    chunk = chunk.strip()
    m = _TITLE_RE.search(chunk)
    if m:
        title = m.group(1).strip()
        body  = chunk[m.end():].strip()
    else:
        body  = chunk
        title = ""
    return title, body


def parse_single_output(text: str):
    return _extract_title_body(text)


def parse_dual_output(text: str):
    main_m = re.search(r'===主稿開始===(.*?)===主稿結束===', text, re.DOTALL)
    side_m = re.search(r'===配稿開始===(.*?)===配稿結束===', text, re.DOTALL)
    mt, mb = _extract_title_body(main_m.group(1)) if main_m else ("", text)
    st, sb = _extract_title_body(side_m.group(1)) if side_m else ("", "")
    return mt, mb, st, sb


# ── file parser ───────────────────────────────────────────────────────────────
def parse_file(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            if not HAS_PDFPLUMBER:
                return "[需安裝 pdfplumber：pip install pdfplumber]"
            import pdfplumber
            out = []
            with pdfplumber.open(path) as pdf:
                for pg in pdf.pages:
                    t = pg.extract_text() or ""
                    if t: out.append(t)
                    for tbl in pg.extract_tables():
                        for row in tbl:
                            if row:
                                out.append("  ".join(str(c) for c in row if c))
            return "\n".join(out)
        elif ext in (".docx", ".doc"):
            if not HAS_DOCX:
                return "[需安裝 python-docx：pip install python-docx]"
            from docx import Document
            return "\n".join(p.text for p in Document(path).paragraphs if p.text)
        elif ext in (".xlsx", ".xls"):
            if not HAS_OPENPYXL:
                return "[需安裝 openpyxl：pip install openpyxl]"
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"[ {ws.title} ]")
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        lines.append("  ".join(str(v) for v in row if v is not None))
            return "\n".join(lines)
        elif ext == ".txt":
            with open(path, encoding="utf-8", errors="replace") as f:
                return f.read()
        else:
            return f"[不支援格式：{ext}]"
    except Exception as e:
        return f"[解析失敗：{e}]"


def strip_markdown(text: str) -> str:
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
    text = re.sub(r'\*{1,3}([^*\n]+)\*{1,3}', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = re.sub(r'```.*?```', '', text, flags=re.DOTALL)
    text = re.sub(r'^[-*_]{3,}\s*$', '', text, flags=re.MULTILINE)
    text = re.sub(r'^\s*[-*+]\s+', '', text, flags=re.MULTILINE)
    text = re.sub(r'^\s*\d+\.\s+', '', text, flags=re.MULTILINE)
    text = re.sub(r'^\s*>\s?', '', text, flags=re.MULTILINE)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def extract_numbers(text: str) -> set:
    return set(re.findall(r"-?\d+(?:,\d{3})*(?:\.\d+)?%?", text))
